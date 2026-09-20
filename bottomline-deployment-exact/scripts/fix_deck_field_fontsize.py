"""One-off patch: add a real `fontFrac` (font size as a fraction of the slide's
own height) to every entry in lib/deckFields.ts, read from the actual pptx run
font size — the previous overlay math derived font size from the *cell's* box
height, which is far taller than a single line of 14pt text, so every edited
figure rendered oversized regardless of bold/colour.

Usage: python scripts/fix_deck_field_fontsize.py
"""

import json
import os
import re

from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIELDS_PATH = os.path.join(ROOT, 'lib', 'deckFields.ts')

DECKS = {
    'jpm': os.path.join(os.environ['TEMP'], 'dealos-artifacts', 'deck-deal-001-53933380132817b8', 'deck.pptx'),
    'illume': os.path.join(os.environ['TEMP'], 'illume_deck.pptx'),
}

ENTRY_RE = re.compile(r'\{[^{}]*\}')
BLOCK_RE = re.compile(r'"(\w+)":\s*\[(.*?)\n  \],', re.DOTALL)


def font_frac_map_for(pptx_path):
    prs = Presentation(pptx_path)
    slide_h = prs.slide_height
    out = {}
    for n, slide in enumerate(prs.slides, start=1):
        for shape_idx, shp in enumerate(slide.shapes):
            if not shp.has_table:
                continue
            table = shp.table
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    sizes = [
                        run.font.size.emu
                        for para in cell.text_frame.paragraphs
                        for run in para.runs
                        if run.font.size
                    ]
                    if sizes:
                        out[f'{n}:{shape_idx}:{ri}:{ci}'] = round(max(sizes) / slide_h, 5)
    return out


def patch_block(fmap, body):
    changed = 0
    def repl(m):
        nonlocal changed
        f = json.loads(m.group(0))
        if f['key'] in fmap:
            f['fontFrac'] = fmap[f['key']]
            changed += 1
        else:
            f.setdefault('fontFrac', 0.022)
        return json.dumps(f)
    return ENTRY_RE.sub(repl, body), changed


def main():
    with open(FIELDS_PATH, 'r', encoding='utf-8') as fh:
        text = fh.read()

    maps = {key: font_frac_map_for(path) for key, path in DECKS.items() if os.path.exists(path)}

    total_changed = 0
    def block_repl(m):
        nonlocal total_changed
        key, body = m.group(1), m.group(2)
        new_body, changed = patch_block(maps.get(key, {}), body)
        total_changed += changed
        print(f'{key}: {changed} fontFrac values set')
        return f'"{key}": [{new_body}\n  ],'

    new_text = BLOCK_RE.sub(block_repl, text)
    new_text = new_text.replace(
        '  bold: boolean;\n}',
        '  bold: boolean;\n  /** Font size as a fraction of the slide\'s own height. */\n  fontFrac: number;\n}',
    )

    with open(FIELDS_PATH, 'w', encoding='utf-8') as fh:
        fh.write(new_text)

    print(f'total: {total_changed} fontFrac values set')


if __name__ == '__main__':
    main()
