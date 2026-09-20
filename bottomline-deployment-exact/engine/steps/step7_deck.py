"""Step 7 — render the client ROI deck from the bank's REAL PowerPoint template.

The template is copied byte-for-byte and only the text runs inside existing
native tables/textboxes and the embedded chart series are overwritten, so every
layout, image, font and piece of branding survives untouched. Nothing is drawn
or recreated.

Data slides and their source tabs in the ROI analysis workbook (File 9):
    9  Payment File Analysis & Conversion Metrics  <- 'Analysis & Conversion'
    10 Your Future State (2 charts + table + copy)  <- 'Payment Transformation'
    11 Strategic Payables may Save You Money        <- 'Cost Savings & Rebate'
    13 Virtual Card Payments                        <- 'Virtual Card'
    14 Premium Network Payments                     <- 'Premium ACH'
    15 Basic Network Payments                       <- 'Basic ACH'
    16 Outsourced Check Payments                    <- 'Outsourced Check Payments'
    17 AP Automation: ROI Summary                   <- 'ROI Summary'
    21 AP Automation Cost Savings Detail            <- 'Cost Savings'

Usage:
    python step7_deck.py --template <real.pptx> --source <File9.xlsx> --out <out.pptx>
"""

import argparse
import json
import shutil
import sys
from decimal import ROUND_HALF_UP, Decimal

import openpyxl
from pptx import Presentation
from pptx.chart.data import CategoryChartData

# 1-based slide number in the deck -> 0-based index used by python-pptx.
DATA_SLIDES = {9: 8, 10: 9, 11: 10, 13: 12, 14: 13, 15: 14, 16: 15, 17: 16, 21: 20}


def log(msg: str) -> None:
    print(msg, file=sys.stderr)


def find_shape(slide, name: str):
    for shp in slide.shapes:
        if shp.name == name:
            return shp
    raise KeyError(f"shape {name!r} not found on slide")


def first_table(slide):
    for shp in slide.shapes:
        if shp.has_table:
            return shp.table
    raise KeyError("no table on slide")


def set_cell_text(cell, text: str) -> None:
    """Overwrite only the first run's text so the original font/colour/bold survives."""
    p = cell.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.text = text


def set_run_text(shape, text: str) -> None:
    shape.text_frame.paragraphs[0].runs[0].text = text


def money(v) -> str:
    return f"${v:,.0f}"


def num(v) -> str:
    return f"{v:,.0f}"


def pct(v) -> str:
    return f"{v * 100:.0f}%"


def mm(v) -> str:
    """Slides 13-16 money convention: exact zero renders as $0.0K, else $X.Xmm."""
    return " $0.0K" if abs(v) < 1 else f" ${v / 1e6:.1f}mm"


def cnt(v) -> str:
    return f"{v:,.0f} "


def build(template: str, source: str, out: str) -> dict:
    shutil.copy(template, out)
    wb = openpyxl.load_workbook(source, data_only=True)
    prs = Presentation(out)
    written = 0

    def fill_product_table(table, ws, row_map):
        """Slides 13-15 share a Spend(E)/Transactions(H)/Vendors(K) row layout."""
        nonlocal written
        for row_idx, src_row in row_map:
            cells = table.rows[row_idx].cells
            set_cell_text(cells[3], mm(ws[f'E{src_row}'].value))
            set_cell_text(cells[6], cnt(ws[f'H{src_row}'].value))
            set_cell_text(cells[9], cnt(ws[f'K{src_row}'].value))
            written += 3

    def fill_split_row(table, ws, row_idx, src_row):
        """The campaign/expected row splits each measure across two columns."""
        nonlocal written
        cells = table.rows[row_idx].cells
        set_cell_text(cells[1], mm(ws[f'C{src_row}'].value))
        set_cell_text(cells[3], mm(ws[f'E{src_row}'].value))
        set_cell_text(cells[4], cnt(ws[f'F{src_row}'].value))
        set_cell_text(cells[6], cnt(ws[f'H{src_row}'].value))
        set_cell_text(cells[7], cnt(ws[f'I{src_row}'].value))
        set_cell_text(cells[9], cnt(ws[f'K{src_row}'].value))
        written += 6

    # --- Slide 9: Payment File Analysis & Conversion Metrics ---
    ac = wb['Analysis & Conversion']
    t9 = first_table(prs.slides[DATA_SLIDES[9]])
    rows9 = [
        ("AP File Totals", 4, money),
        ("LESS: Exclusions", 5, money),
        ("Electronic Payment Targets", 6, money),
        ("Virtual Card Conversion", 7, money),
        ("Premium ACH Conversion", 8, money),
        ("Basic ACH Conversion", 9, money),
        ("Card & ACH Conversion %", 10, pct),
        ("Paymode Network Matches", 11, money),
        ("Paymode Network Matches %", 12, pct),
    ]
    for i, (label, r, spend_fmt) in enumerate(rows9, start=1):
        cells = t9.rows[i].cells
        count_fmt = pct if spend_fmt is pct else num
        set_cell_text(cells[0], label)
        set_cell_text(cells[1], count_fmt(ac[f'C{r}'].value))
        set_cell_text(cells[2], count_fmt(ac[f'D{r}'].value))
        set_cell_text(cells[3], spend_fmt(ac[f'E{r}'].value))
        written += 4

    # --- Slide 10: Your Future State ---
    s10 = prs.slides[DATA_SLIDES[10]]
    pt = wb['Payment Transformation']
    set_run_text(find_shape(s10, 'TextBox 4'), pt['B5'].value)

    # Current state has no Paymode-X ACH volume yet, hence the hard 0 in slot 2.
    current = [pt['C16'].value, 0, pt['C18'].value, pt['C19'].value, pt['C20'].value, pt['C21'].value]
    cd = CategoryChartData()
    cd.categories = ['Virtual Card', 'Paymode-X ACH', 'Check', 'ACH', 'Wire', 'Card']
    cd.add_series('Values', [c / pt['C22'].value for c in current])
    cd.add_series('Column1', current)
    find_shape(s10, 'Chart 7').chart.replace_data(cd)

    # Counts must be rounded before the share is derived, to match the real deck.
    future = [round(pt[f'F{r}'].value) for r in range(16, 22)]
    cd2 = CategoryChartData()
    cd2.categories = ['Virtual Card', 'Paymode ACH', 'Outsourced Check', 'ACH', 'Wire', 'Card']
    cd2.add_series('Values', [c / pt['F22'].value for c in future])
    cd2.add_series('Column1', future)
    find_shape(s10, 'Chart 14').chart.replace_data(cd2)

    t10 = first_table(s10)
    keys = ['Virtual Card', 'Paymode ACH', 'Check', 'ACH', 'Wire', 'Card', 'Total']
    future_labels = dict(zip(keys, ['Virtual Card', 'Paymode ACH', 'Outsourced Check', 'ACH', 'Wire', 'Card', 'Total']))
    for i, key in enumerate(keys):
        src = 16 + i
        cells = t10.rows[i].cells
        set_cell_text(cells[0], key)
        set_cell_text(cells[1], num(pt[f'C{src}'].value))
        set_cell_text(cells[3], future_labels[key])
        set_cell_text(cells[4], num(pt[f'F{src}'].value))
        written += 4

    # --- Slide 11: Strategic Payables may Save You Money ---
    csr = wb['Cost Savings & Rebate']
    t11 = first_table(prs.slides[DATA_SLIDES[11]])
    set_cell_text(t11.rows[0].cells[0], money(csr['B6'].value) + " ")
    set_cell_text(t11.rows[4].cells[0], money(csr['B10'].value) + " ")
    written += 2

    # --- Slide 13: Virtual Card Payments ---
    s13 = prs.slides[DATA_SLIDES[13]]
    vc = wb['Virtual Card']
    t13 = first_table(s13)
    fill_product_table(t13, vc, [(1, 3), (2, 4), (4, 6)])
    fill_split_row(t13, vc, 3, 5)
    set_run_text(find_shape(s13, 'TextBox 11'), f" ${vc['M10'].value / 1000:.1f}k")
    set_run_text(find_shape(s13, 'TextBox 16'), f" ${vc['M11'].value / 1000:.1f}k")
    written += 2

    # --- Slide 14: Premium Network Payments ---
    pa = wb['Premium ACH']
    t14 = first_table(prs.slides[DATA_SLIDES[14]])
    fill_product_table(t14, pa, [(1, 3), (3, 5)])
    fill_split_row(t14, pa, 2, 4)

    # --- Slide 15: Basic Network Payments ---
    s15 = prs.slides[DATA_SLIDES[15]]
    ba = wb['Basic ACH']
    t15 = first_table(s15)
    fill_product_table(t15, ba, [(1, 3), (3, 5)])
    fill_split_row(t15, ba, 2, 4)
    set_run_text(find_shape(s15, 'TextBox 34'), f"(${ba['M11'].value / 1000:.1f}k)")
    written += 1

    # --- Slide 16: Outsourced Check Payments ---
    ocp = wb['Outsourced Check Payments']
    t16 = first_table(prs.slides[DATA_SLIDES[16]])
    for i, (label, r) in enumerate([("Check Print", 6), (" Card & ACH Declines", 7), ("Total", 8)], start=1):
        cells = t16.rows[i].cells
        set_cell_text(cells[0], label)
        set_cell_text(cells[1], mm(ocp[f'C{r}'].value))
        set_cell_text(cells[2], cnt(ocp[f'D{r}'].value))
        set_cell_text(cells[3], cnt(ocp[f'E{r}'].value))
        written += 4

    # --- Slide 17: AP Automation ROI Summary ---
    roi = wb['ROI Summary']
    t17 = first_table(prs.slides[DATA_SLIDES[17]])
    roi_labels = ["Virtual Card Payments", "Premium Network Payments", "Basic Network Payments",
                  "Outsourced Check Payments", "TOTALS"]
    for i, label in enumerate(roi_labels):
        r = 6 + i
        cells = t17.rows[i + 1].cells
        set_cell_text(cells[0], label)
        for cell_idx, col in zip((1, 3, 5, 7, 8), 'CEGIJ'):
            set_cell_text(cells[cell_idx], num(roi[f'{col}{r}'].value))
        written += 6

    # --- Slide 21: Cost Savings Detail Summary ---
    s21 = prs.slides[DATA_SLIDES[21]]
    cs = wb['Cost Savings']
    t21 = first_table(s21)
    for i, label in enumerate(["Virtual Card", "Premium ACH", "Basic ACH", "Outsourced Check", "Total"]):
        r = 6 + i
        cells = t21.rows[i + 1].cells
        # These cells use Excel's accounting format: value (or a dash) padded with a
        # trailing currency gutter, then right-justified to the template's own width.
        orig_widths = [len(c.text) for c in cells]
        set_cell_text(cells[0], label)
        for j, col in enumerate('CDEF', start=1):
            v = cs[f'{col}{r}'].value
            txt = f"{v:,.0f} " if v != 0 else "-   "
            set_cell_text(cells[j], txt.rjust(max(len(txt), orig_widths[j])))
        set_cell_text(cells[5], f"${cs[f'G{r}'].value:,.0f}")
        written += 6
    set_run_text(
        find_shape(s21, 'TextBox 35'),
        f"Cost savings will be ${cs['G10'].value / 1000:.1f}k assuming current transaction costs are:",
    )
    written += 1

    prs.save(out)
    return {
        'status': 'ok',
        'outputPath': out,
        'slideCount': len(prs.slides),
        'dataSlides': sorted(DATA_SLIDES),
        'valuesWritten': written,
    }


def build_illume(template: str, source: str, out: str) -> dict:
    """Fifth Third / Illume Ag deck — 19 slides, a different layout to the JPM one.

    Every product slide here uses the same 4-column table (label, Spend,
    Transactions, Vendors) fed from a sheet whose data sits in B=label,
    C/D/E=measures, so the whole deck is described by the table below rather
    than by per-slide code.

    Callout textboxes are deliberately not written: they are all named
    'Title 1', so addressing them would depend on shape order, and a wrong
    guess would put a figure in the wrong box on a client-facing deck.
    """
    shutil.copy(template, out)
    wb = openpyxl.load_workbook(source, data_only=True)
    prs = Presentation(out)
    written = 0

    def spend(v) -> str:
        if v is None:
            return ''
        return f"${v / 1e6:.1f}mm" if abs(v) >= 1e6 else f"${v / 1e3:.1f}K"

    def xnum(v) -> str:
        """Excel rounds .5 away from zero; Python's round() goes to even, which
        turns 34.5 into 34 and silently disagrees with the delivered deck."""
        if v is None:
            return ''
        return f"{Decimal(str(v)).quantize(Decimal('1'), rounding=ROUND_HALF_UP):,}"

    # slide (1-based) -> (sheet, first data row, number of rows)
    product_slides = [
        (7, 'Virtual Card', 3, 4),
        (8, 'Premium ACH', 8, 4),
        (9, 'Basic ACH', 3, 4),
        (10, 'ACH Migration', 3, 3),
        (11, 'Outsourced Check Payments', 3, 3),
        (12, 'B2C', 3, 3),
    ]

    slides_touched = []

    # --- Slide 6: Payment File Analysis & Conversion Metrics ---
    ac = wb['Analysis & Conversion']
    t13 = first_table(prs.slides[5])
    # ($ on the total/target/match rows only, matching the delivered deck)
    rows13 = [(3, True), (4, False), (5, True), (6, False), (7, False),
              (8, None), (9, True), (10, None)]
    for i, (r, dollar) in enumerate(rows13, start=1):
        cells = t13.rows[i].cells
        if dollar is None:  # percentage row
            set_cell_text(cells[1], pct(ac[f'C{r}'].value))
            set_cell_text(cells[2], pct(ac[f'D{r}'].value))
            set_cell_text(cells[3], pct(ac[f'E{r}'].value))
        else:
            set_cell_text(cells[1], xnum(ac[f'C{r}'].value))
            set_cell_text(cells[2], xnum(ac[f'D{r}'].value))
            e = ac[f'E{r}'].value
            set_cell_text(cells[3], f"${xnum(e)}" if dollar else xnum(e))
        written += 3
    slides_touched.append(6)

    # --- Product tables ---
    for slide_no, sheet_name, first_row, n_rows in product_slides:
        ws = wb[sheet_name]
        table = first_table(prs.slides[slide_no - 1])
        for i in range(n_rows):
            r = first_row + i
            cells = table.rows[i + 1].cells
            set_cell_text(cells[1], spend(ws[f'C{r}'].value))
            set_cell_text(cells[2], xnum(ws[f'D{r}'].value))
            set_cell_text(cells[3], xnum(ws[f'E{r}'].value))
            written += 3
        slides_touched.append(slide_no)

    # The reworked deck dropped the standalone Invoice Automation slide; its
    # figures now only appear in the ROI and cost-savings summaries.

    # --- Slide 17: ROI Summary ---
    rs = wb['ROI Summary']
    t25 = first_table(prs.slides[16])
    for i, r in enumerate(range(3, 11), start=1):
        cells = t25.rows[i].cells
        for col, letter in ((1, 'C'), (3, 'E'), (5, 'G'), (7, 'I'), (8, 'J')):
            set_cell_text(cells[col], xnum(rs[f'{letter}{r}'].value))
            written += 1
    slides_touched.append(17)

    # --- Slide 19: Cost Savings detail ---
    cs = wb['Cost Savings']
    t27 = first_table(prs.slides[18])
    for i, r in enumerate(range(3, 11), start=1):
        cells = t27.rows[i].cells
        for col, letter in ((1, 'C'), (2, 'D'), (3, 'E'), (4, 'F')):
            v = cs[f'{letter}{r}'].value
            # the invoice-automation row carries a label instead of a count
            set_cell_text(cells[col], xnum(v) if isinstance(v, (int, float)) else str(v or ''))
            written += 1
        set_cell_text(cells[5], f"${xnum(cs[f'G{r}'].value)}")
        written += 1
    slides_touched.append(19)

    prs.save(out)
    return {
        'status': 'ok',
        'outputPath': out,
        'slideCount': len(prs.slides),
        'dataSlides': slides_touched,
        'valuesWritten': written,
    }


BUILDERS = {'jpm': build, 'illume': build_illume}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument('--template', required=True, help='the bank\'s real .pptx deck')
    ap.add_argument('--source', required=True, help='File 9 ROI analysis workbook')
    ap.add_argument('--out', required=True)
    ap.add_argument('--map', default='jpm', choices=sorted(BUILDERS),
                    help='which deck layout the template follows')
    a = ap.parse_args()

    log(f"copying template {a.template}")
    result = BUILDERS[a.map](a.template, a.source, a.out)
    log(f"wrote {result['valuesWritten']} values across {len(result['dataSlides'])} data slides")
    print(json.dumps(result))


if __name__ == '__main__':
    main()
