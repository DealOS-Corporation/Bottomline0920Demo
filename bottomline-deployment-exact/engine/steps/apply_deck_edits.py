"""Apply analyst edits to a generated deck.

Keys are the same `slide:shape:row:col` ids emitted by scripts/deck_fields.py, so
whatever the analyst typed over a figure in the preview is what lands in the
downloaded .pptx. Formatting is preserved by reusing step7_deck's run-level
setter rather than assigning cell.text.

Usage:
    python apply_deck_edits.py --deck <generated.pptx> --edits <edits.json>
"""

import argparse
import json

from pptx import Presentation

from step7_deck import set_cell_text


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument('--deck', required=True)
    ap.add_argument('--edits', required=True, help='JSON file: {"slide:shape:row:col": "text"}')
    a = ap.parse_args()

    with open(a.edits, encoding='utf-8') as fh:
        edits = json.load(fh)

    prs = Presentation(a.deck)
    applied, skipped = 0, []

    for key, value in edits.items():
        try:
            slide_no, shape_idx, row, col = (int(p) for p in key.split(':'))
        except ValueError:
            skipped.append(key)
            continue
        if not (1 <= slide_no <= len(prs.slides)):
            skipped.append(key)
            continue
        shapes = list(prs.slides[slide_no - 1].shapes)
        if shape_idx >= len(shapes) or not shapes[shape_idx].has_table:
            skipped.append(key)
            continue
        table = shapes[shape_idx].table
        if row >= len(table.rows) or col >= len(table.columns):
            skipped.append(key)
            continue
        set_cell_text(table.rows[row].cells[col], str(value))
        applied += 1

    prs.save(a.deck)
    print(json.dumps({'status': 'ok', 'applied': applied, 'skipped': skipped}))


if __name__ == '__main__':
    main()
