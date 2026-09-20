"""Step 5 — Classification Review engine.

CLI:
    python step5_classify.py --in <File6.xlsx> --out <File7_generated.xlsx>

Reads the Alteryx pricing output (File 6), applies the industry override
rules, writes the reviewed output (File 7 shape: same columns + 'Segment
(Original)' + 'Override Reason', changed rows sorted first), and prints a
single JSON object (StepResult) to stdout for the Node caller to parse.
Never prints anything else to stdout — logs/errors go to stderr.
"""
import argparse
import json
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from rules_industry import classify_override  # noqa: E402


def load_rows(path: str):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    raw = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() if h is not None else "" for h in raw[0]]
    rows = []
    for r in raw[1:]:
        if all(c is None for c in r):
            continue
        rows.append(dict(zip(headers, r)))
    return headers, rows


def run(in_path: str, out_path: str):
    headers, rows = load_rows(in_path)
    changed = []
    unchanged = []

    for row in rows:
        segment = row.get("Segment")
        new_segment, reason = classify_override(
            row.get("Vendor Name"), row.get("Country"), segment
        )
        out_row = dict(row)
        if new_segment:
            out_row["Segment (Original)"] = segment
            out_row["Override Reason"] = reason
            out_row["Segment"] = new_segment
            changed.append(out_row)
        else:
            out_row["Segment (Original)"] = ""
            out_row["Override Reason"] = ""
            unchanged.append(out_row)

    out_headers = headers + ["Segment (Original)", "Override Reason"]
    ordered_rows = changed + unchanged

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "sheet0"
    ws_out.append(out_headers)
    for row in ordered_rows:
        ws_out.append([row.get(h, "") for h in out_headers])
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb_out.save(out_path)

    total = len(rows)
    result = {
        "status": "done",
        "metrics": [
            {"label": "Vendors reviewed", "value": f"{total:,}"},
            {"label": "Segment overrides", "value": str(len(changed)), "sub": f"{len(changed) / total * 100:.1f}% of vendors"},
        ],
        "checks": [
            {
                "id": "override-rules",
                "label": "Segment override rules",
                "status": "warn" if changed else "pass",
                "detail": f"{len(changed)} of {total} vendors had their Segment changed by industry override rules.",
            }
        ],
        "overrides": [
            {
                "vendor": r.get("Vendor Name"),
                "segmentBefore": r.get("Segment (Original)"),
                "segmentAfter": r.get("Segment"),
                "reason": r.get("Override Reason"),
            }
            for r in changed
        ],
        "outputPath": out_path,
    }
    print(json.dumps(result))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--in", dest="in_path", required=True)
    parser.add_argument("--out", dest="out_path", required=True)
    args = parser.parse_args()
    try:
        run(args.in_path, args.out_path)
    except Exception as exc:  # surfaced to Node as stderr + non-zero exit
        print(f"step5_classify failed: {exc}", file=sys.stderr)
        sys.exit(1)
