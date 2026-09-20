"""Verify the step5 classification engine against the real, human-reviewed
File 7 for the JPM/Denton Co Electric Coop deal.

Usage: python verify_step5.py
Prints hit / miss / false-positive rows to stdout, exit code 1 if any
false positives are found (those are worse than misses — they mean the
engine changed a row a human did NOT change).
"""
import json
import sys
from pathlib import Path

import openpyxl

ENGINE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ENGINE_DIR))
from steps.step5_classify import run as run_step5  # noqa: E402

JPM_DIR = Path(
    r"C:\Users\91914\OneDrive\Desktop\Company Files\Customer\Bottomline\POC Deals\JPM"
)
FILE6 = JPM_DIR / "6_Alteryx Pricing Workflow Output_DCEC_PricingOutput.xlsx"
FILE7_REAL = JPM_DIR / "7_Manually Reviewed Alteryx Output_DCEC_PricingOutput_MR.xlsx"
OUT = Path(r"C:\Users\91914\AppData\Local\Temp\bt_engine_probe\7_generated.xlsx")


def load_by_id(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = {}
    for r in rows[1:]:
        if all(c is None for c in r):
            continue
        d = dict(zip(headers, r))
        out[str(d.get("ROC Vendor ID"))] = d
    return out


def main():
    run_step5(str(FILE6), str(OUT))

    real6 = load_by_id(FILE6)
    real7 = load_by_id(FILE7_REAL)
    generated7 = load_by_id(OUT)

    real_changed = {k for k in real7 if real6.get(k, {}).get("Segment") != real7[k].get("Segment")}
    engine_changed = {k for k in generated7 if real6.get(k, {}).get("Segment") != generated7[k].get("Segment")}

    hits = real_changed & engine_changed
    misses = real_changed - engine_changed
    false_positives = engine_changed - real_changed

    print(f"Real human-reviewed changes: {len(real_changed)}")
    print(f"Engine changes:              {len(engine_changed)}")
    print(f"Hits:                        {len(hits)}")
    print(f"Misses (engine didn't catch):{len(misses)}")
    print(f"False positives (engine changed, human didn't):{len(false_positives)}")

    def describe(ids, source):
        for k in ids:
            row = source.get(k, {})
            print(f"  [{k}] {row.get('Vendor Name')!r} country={row.get('Country')} "
                  f"before={real6.get(k, {}).get('Segment')!r} -> {row.get('Segment')!r}")

    if misses:
        print("\nMISSES:")
        describe(misses, real7)
    if false_positives:
        print("\nFALSE POSITIVES:")
        describe(false_positives, generated7)
    if hits:
        print("\nHITS:")
        describe(hits, generated7)

    sys.exit(1 if false_positives else 0)


if __name__ == "__main__":
    main()
